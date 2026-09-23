"""Pruebas de la primera iteración de Reporte de Evidencias Mobile."""

from hashlib import sha256
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

from docx import Document

from mobile_evidence_generator import (
    MobileEvidenceData,
    generate_mobile_evidence,
    generate_mobile_report,
    mobile_evidence_filename,
    parse_txt,
    validate_report_data,
)
from resource_paths import resource_path
from evidence_generator import EvidenceCase


class MobileEvidenceGeneratorTests(unittest.TestCase):
    def setUp(self) -> None:
        text = Path("tests/data/evidencia_mobile_example.txt").read_text(encoding="utf-8")
        self.data = MobileEvidenceData.from_mapping(parse_txt(text))

    def test_parse_txt_handles_spacing_utf8_and_first_equals_only(self) -> None:
        values = parse_txt("\n QA_AGILE = Diego Muñoz \nCASO_PRUEBA=Valor=A=B\n")
        self.assertEqual("Diego Muñoz", values["QA_AGILE"])
        self.assertEqual("Valor=A=B", values["CASO_PRUEBA"])

    def test_parse_txt_reports_invalid_line(self) -> None:
        with self.assertRaisesRegex(ValueError, "Línea 2 inválida"):
            parse_txt("QA_AGILE=Diego\nsin separador")

    def test_validate_report_data(self) -> None:
        invalid = MobileEvidenceData(status="DONE", execution_date="2026-09-23")
        errors = validate_report_data(invalid)
        self.assertIn("El campo QA Agile es obligatorio", errors)
        self.assertIn("Fecha de ejecución inválida. Utilice DD/MM/YYYY", errors)
        self.assertIn("Estado inválido. Estados permitidos: PASSED, FAILED, BLOCKED, NOT_EXECUTED", errors)

    def test_generates_header_without_modifying_template(self) -> None:
        template = resource_path("templates/Evidencias_Mobile_template.docx")
        original_hash = sha256(template.read_bytes()).hexdigest()
        with TemporaryDirectory() as temporary_directory:
            output = Path(temporary_directory) / mobile_evidence_filename("URPIPRO-123")
            generate_mobile_report(self.data, template, output)
            generated_text = "\n".join(cell.text for table in Document(output).tables for row in table.rows for cell in row.cells)
            for expected in (
                "Diego Molina",
                "23/09/2026",
                "Validación E2E Mobile - Simulación de crédito",
                "PASSED",
                "QA2 - Android - App 2.0.0",
            ):
                self.assertIn(expected, generated_text)
        self.assertEqual(original_hash, sha256(template.read_bytes()).hexdigest())
        self.assertEqual("Evidencia_Mobile_ADN-234.docx", mobile_evidence_filename("ADN-234"))

    def test_generates_a_cloned_scenario_for_each_detected_case(self) -> None:
        template = resource_path("templates/Evidencias_Mobile_template.docx")
        original_hash = sha256(template.read_bytes()).hexdigest()
        case_path = "Urpipro/Loan/Derivacion App MiBanco"
        cases = [
            EvidenceCase("TAP-78824", f"{case_path}/Visualización correcta del modal de confirmación"),
            EvidenceCase("TAP-78825", f"{case_path}/Confirmación y procesamiento de la solicitud"),
            EvidenceCase("TAP-78826", f"{case_path}/Cancelación y cierre del modal"),
            EvidenceCase("TAP-78827", f"{case_path}/Validar simulación con monto máximo"),
        ]
        data = MobileEvidenceData(qa_agile="Diego", execution_date="23/09/2026", test_plan="Pruebas E2E Mobile")

        with TemporaryDirectory() as temporary_directory:
            output = Path(temporary_directory) / mobile_evidence_filename("ADN-234")
            generate_mobile_evidence(data, cases, template, output)
            document = Document(output)
            text = "\n".join(
                [paragraph.text for paragraph in document.paragraphs]
                + [cell.text for table in document.tables for row in table.rows for cell in row.cells]
            )
            self.assertEqual(5, len(document.tables))
            self.assertEqual([(3, 2)] * 4, [(len(table.rows), len(table.columns)) for table in document.tables[1:]])
            self.assertEqual(case_path, document.tables[0].cell(1, 1).text)
            self.assertEqual([f"Escenario de prueba {number}: {case.title.rsplit('/', 1)[1]}" for number, case in enumerate(cases, start=1)], [
                paragraph.text for paragraph in document.paragraphs if paragraph.text.startswith("Escenario de prueba ")
            ])
            self.assertIn("Diego", text)
            self.assertIn("Pruebas E2E Mobile", text)
            for table, case in zip(document.tables[1:], cases):
                self.assertEqual(f"Caso: {case.xray_test}", table.cell(0, 0).text)
                self.assertEqual("Estado: Validado", table.cell(0, 1).text)
                self.assertEqual("", table.cell(1, 0).text)
        self.assertEqual(original_hash, sha256(template.read_bytes()).hexdigest())

    def test_mobile_header_uses_selected_environment_and_fixed_expected_result(self) -> None:
        template = resource_path("templates/Evidencias_Mobile_template.docx")
        cases = [EvidenceCase("TAP-78744", "Urpipro/Loan/Derivacion App MiBanco/Visualización correcta")]
        expected_result = (
            "Que se cumplan todos los escenarios planificados. Todos los escenarios deben finalizar satisfactoriamente y cumplir con lo esperado."
        )

        for environment in ("QA", "STG"):
            with self.subTest(environment=environment), TemporaryDirectory() as temporary_directory:
                data = MobileEvidenceData(
                    qa_agile="Diego", execution_date="23/09/2026", test_plan="5678", test_environment=environment
                )
                output = Path(temporary_directory) / mobile_evidence_filename("ADN-234")
                generate_mobile_evidence(data, cases, template, output)
                header = Document(output).tables[0]
                self.assertEqual(environment, header.cell(2, 3).text)
                self.assertEqual(expected_result, header.cell(3, 1).text)


if __name__ == "__main__":
    unittest.main()
