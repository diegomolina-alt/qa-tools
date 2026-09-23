"""Parsing y generación de evidencias API desde la plantilla oficial DOCX."""

from __future__ import annotations

from copy import deepcopy
from dataclasses import dataclass
from pathlib import Path
import re

from docx import Document
from docx.oxml import OxmlElement
from docx.table import Table
from docx.text.paragraph import Paragraph

from docx_helpers import set_scenario_title_text
from filename_helpers import sanitize_filename_component


XRAY_TEST_PATTERN = re.compile(r"TAP-\d+", re.IGNORECASE)
ACCEPTANCE_CRITERION_PATTERN = re.compile(r"^\s*(?P<criterion>CA\s*\d+)\s*:\s*(?P<title>.*)$", re.IGNORECASE)
MARKDOWN_LINK_PATTERN = re.compile(r"^\[([^\]]+)\]\([^)]*\)$")
MARKDOWN_SEPARATOR_PATTERN = re.compile(r"^:?-+:?$")
DEFAULT_DOCUMENT_STATUS = "PASSED"
XRAY_METADATA = frozenset(
    {
        "xray test",
        "cucumber",
        "0",
        "to do",
        "pass",
        "failed",
        "fail",
        "executing",
        "aborted",
        "blocked",
    }
)


@dataclass
class EvidenceCase:
    xray_test: str
    title: str
    method: str = "POST"
    endpoint: str = ""
    status_response: str = ""
    value_status_response: str = ""
    request: str = ""
    response: str = ""
    verification: str = ""
    acceptance_criterion: str = ""
    content: str = ""

    @property
    def display_title(self) -> str:
        """Título para la lista, incluyendo el identificador CA cuando existe."""
        return f"{self.acceptance_criterion}: {self.title}" if self.acceptance_criterion else self.title

@dataclass(frozen=True)
class ApiEvidenceData:
    jira: str
    qa_agile: str
    execution_date: str
    service: str
    test_plan: str
    cases: list[EvidenceCase]
    test_environment: str = ""


def parse_xray_tests(text: str) -> list[EvidenceCase]:
    """Extrae TAP/título de los formatos de Xray y Test Executions, en su orden original."""
    markdown_cases = _parse_markdown_xray_tests(text)
    if markdown_cases:
        return _unique_cases(markdown_cases)

    # Jira puede copiar una fila de tabla con columnas separadas por tabuladores.
    # Convertirlas en elementos lógicos permite usar el mismo flujo que con líneas.
    lines = [line.strip() for line in re.split(r"[\r\n\t]+", text)]
    cases: list[EvidenceCase] = []
    tap_indexes = [index for index, line in enumerate(lines) if XRAY_TEST_PATTERN.fullmatch(line)]

    for position, tap_index in enumerate(tap_indexes):
        next_tap_index = tap_indexes[position + 1] if position + 1 < len(tap_indexes) else len(lines)
        title = _first_title_line(lines, tap_index + 1, next_tap_index)
        if title:
            cases.append(EvidenceCase(lines[tap_index].upper(), title))

    return _unique_cases(cases)


def parse_test_titles(text: str) -> list[EvidenceCase]:
    """Interpreta títulos copiados de Jira y conserva el formato Xray cuando existe.

    Jira puede entregar criterios de aceptación. Cada bloque que inicia con
    CAxx: representa un caso completo, incluidas sus líneas Dado/Cuando/Entonces.
    Si no hay TAP ni criterios, se mantiene el soporte para listas simples.
    """
    cases = parse_xray_tests(text)
    if cases:
        return cases
    criteria = parse_acceptance_criteria(text)
    if criteria:
        return criteria
    titles = [line.strip() for line in text.splitlines() if line.strip()]
    return [EvidenceCase(f"CASO-{number:03d}", title) for number, title in enumerate(titles, start=1)]


def parse_acceptance_criteria(text: str) -> list[EvidenceCase]:
    """Agrupa cada CAxx: y todo su contenido hasta el siguiente encabezado CA."""
    cases: list[EvidenceCase] = []
    criterion = ""
    title = ""
    content_lines: list[str] = []

    def add_current() -> None:
        if criterion:
            cases.append(
                EvidenceCase(
                    xray_test=f"CASO-{len(cases) + 1:03d}",
                    title=title,
                    acceptance_criterion=criterion,
                    content="\n".join(content_lines).strip(),
                )
            )

    for line in text.splitlines():
        match = ACCEPTANCE_CRITERION_PATTERN.match(line)
        if match:
            add_current()
            criterion = re.sub(r"\s+", "", match.group("criterion")).upper()
            title = match.group("title").strip()
            content_lines = []
        elif criterion:
            content_lines.append(line)

    add_current()
    return cases


def generate_api_evidence(data: ApiEvidenceData, template_path: Path, output_path: Path) -> None:
    """Copia la plantilla y genera un bloque DOCX independiente por cada caso."""
    if not data.cases:
        raise ValueError("Debe existir al menos un caso de evidencia.")

    document = _load_api_template(template_path)
    general_table, template_evidence_table = document.tables
    template_heading = template_evidence_table._tbl.getprevious()
    domain = ""
    evidence_names: list[str] = []
    for case in data.cases:
        detected_domain, evidence_name = _split_api_title(case.title)
        if detected_domain and not domain:
            domain = detected_domain
        evidence_names.append(evidence_name)

    _set_cell_text(general_table.cell(0, 1), data.qa_agile)
    _set_cell_text(general_table.cell(0, 3), data.execution_date)
    _set_cell_text(general_table.cell(1, 1), data.service)
    _set_cell_text(general_table.cell(2, 1), domain)
    _set_cell_text(general_table.cell(3, 1), DEFAULT_DOCUMENT_STATUS)
    _set_cell_text(general_table.cell(3, 3), data.test_environment)
    _set_cell_text(general_table.cell(4, 1), data.test_plan)

    evidence_blocks = [(template_heading, template_evidence_table)]
    for _ in data.cases[1:]:
        evidence_blocks.append(_duplicate_evidence_block(evidence_blocks[-1][1], document, template_heading))

    for number, (case, evidence_name, (heading, table)) in enumerate(
        zip(data.cases, evidence_names, evidence_blocks), start=1
    ):
        set_scenario_title_text(
            Paragraph(heading, document._body), f"Escenario de prueba {number}: {evidence_name}"
        )
        _set_cell_text(table.cell(0, 1), case.xray_test)
        _set_cell_text(table.cell(0, 3), case.method)
        _set_cell_text(table.cell(1, 1), case.endpoint)
        _set_cell_text(table.cell(2, 1), case.status_response)
        _set_cell_text(table.cell(2, 3), case.value_status_response)
        _set_cell_text(table.cell(3, 1), case.request)
        _set_cell_text(table.cell(4, 1), case.response)
        _set_cell_text(table.cell(5, 1), case.verification)

    document.save(output_path)


def evidence_filename(jira: str) -> str:
    """Construye un nombre de archivo seguro sin modificar guiones válidos."""
    safe_jira = sanitize_filename_component(jira, fallback="SIN_JIRA")
    return f"Evidencias_APIs_{safe_jira}.docx"


def _load_api_template(template_path: Path) -> Document:
    """Carga la plantilla API oficial y verifica su cabecera y bloque de escenario."""
    document = Document(template_path)
    if len(document.tables) < 2:
        raise ValueError("La plantilla API no contiene la cabecera y el bloque de escenario requeridos.")

    expected_labels = {"QA Agile:", "Servicio:", "Dominio:", "Estado:", "Entorno de prueba:", "Test Plan:"}
    actual_labels = {cell.text.strip() for row in document.tables[0].rows for cell in row.cells}
    if not expected_labels.issubset(actual_labels):
        raise ValueError("La plantilla API no contiene las etiquetas de cabecera requeridas.")

    scenario_heading = document.tables[1]._tbl.getprevious()
    if scenario_heading is None or "Escenario de prueba:" not in Paragraph(scenario_heading, document._body).text:
        raise ValueError("La plantilla API no contiene el bloque 'Escenario de prueba:' requerido.")
    return document


def _split_api_title(title: str) -> tuple[str, str]:
    """Separa Dominio y Evidencia por la última barra sin alterar rutas internas."""
    if "/" not in title:
        return "", title.strip()
    domain, evidence_name = title.rsplit("/", 1)
    return domain.strip(), evidence_name.strip()


def _first_title_line(lines: list[str], start: int, end: int) -> str | None:
    """Obtiene el primer valor útil de la fila/bloque que sigue a un TAP."""
    for line in lines[start:end]:
        normalized = line.strip()
        if normalized and not _is_xray_metadata(normalized):
            return normalized
    return None


def _is_xray_metadata(value: str) -> bool:
    """Identifica columnas auxiliares de Jira/Xray que no pueden ser un título."""
    return value.casefold() in XRAY_METADATA or value.isdecimal()


def _parse_markdown_xray_tests(text: str) -> list[EvidenceCase]:
    """Lee filas Markdown de Jira sin depender de la posición de sus columnas."""
    cases: list[EvidenceCase] = []
    for row in text.splitlines():
        cells = _markdown_cells(row)
        if not cells or _is_markdown_separator(cells):
            continue

        tap_index, xray_test = _find_tap_cell(cells)
        if tap_index is None or xray_test is None:
            continue

        title = _first_markdown_title(cells, tap_index + 1)
        if title:
            cases.append(EvidenceCase(xray_test, title))
    return cases


def _markdown_cells(row: str) -> list[str] | None:
    """Devuelve las celdas de una fila de tabla Markdown, incluidas las vacías."""
    stripped = row.strip()
    if "|" not in stripped:
        return None
    return [_clean_markdown_cell(cell) for cell in stripped.strip("|").split("|")]


def _is_markdown_separator(cells: list[str]) -> bool:
    meaningful_cells = [cell for cell in cells if cell]
    return bool(meaningful_cells) and all(MARKDOWN_SEPARATOR_PATTERN.fullmatch(cell) for cell in meaningful_cells)


def _find_tap_cell(cells: list[str]) -> tuple[int | None, str | None]:
    for index, cell in enumerate(cells):
        match = XRAY_TEST_PATTERN.search(cell)
        if match:
            return index, match.group(0).upper()
    return None, None


def _first_markdown_title(cells: list[str], start: int) -> str | None:
    for cell in cells[start:]:
        if not cell or _is_xray_metadata(cell) or cell.startswith(("http://", "https://")):
            continue
        # Una celda con un enlace a otro TAP o a Jira es metadata, no un resumen.
        if MARKDOWN_LINK_PATTERN.fullmatch(cell) or XRAY_TEST_PATTERN.fullmatch(cell):
            continue
        return cell
    return None


def _clean_markdown_cell(cell: str) -> str:
    """Quita únicamente sintaxis de presentación Markdown de una celda."""
    value = cell.strip()
    link = MARKDOWN_LINK_PATTERN.fullmatch(value)
    if link:
        value = link.group(1).strip()
    return value.replace("**", "").strip()


def _unique_cases(cases: list[EvidenceCase]) -> list[EvidenceCase]:
    """Conserva la primera aparición de cada TAP, respetando el orden pegado."""
    seen_tests: set[str] = set()
    unique_cases: list[EvidenceCase] = []
    for case in cases:
        if case.xray_test not in seen_tests:
            seen_tests.add(case.xray_test)
            unique_cases.append(case)
    return unique_cases


def _duplicate_evidence_block(table: Table, document: Document, template_heading: object) -> tuple[object, Table]:
    """Duplica el rótulo y la tabla desde el XML de la plantilla original."""
    separator = OxmlElement("w:p")
    copied_heading = deepcopy(template_heading)
    copied_table = deepcopy(table._tbl)
    table._tbl.addnext(separator)
    separator.addnext(copied_heading)
    copied_heading.addnext(copied_table)
    return copied_heading, Table(copied_table, document._body)


def _set_cell_text(cell: object, value: str) -> None:
    paragraphs = cell.paragraphs
    _set_paragraph_text(paragraphs[0], value)
    for paragraph in paragraphs[1:]:
        _set_paragraph_text(paragraph, "")


def _set_paragraph_text(paragraph: Paragraph, value: str) -> None:
    for run in paragraph.runs:
        run.text = ""
    paragraph.add_run(value)
