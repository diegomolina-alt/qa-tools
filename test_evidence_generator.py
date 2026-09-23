"""Pruebas de parsing y generación de evidencias API por Xray Test."""

from hashlib import sha256
from pathlib import Path
from tempfile import TemporaryDirectory
from zipfile import ZipFile
import unittest

from docx import Document

from evidence_generator import DEFAULT_DOCUMENT_STATUS, ApiEvidenceData, EvidenceCase, evidence_filename, generate_api_evidence, parse_acceptance_criteria, parse_test_titles, parse_xray_tests
from resource_paths import resource_path


def _scenario_title_format(paragraph: object) -> tuple[object, ...]:
    """Propiedades visuales que API debe conservar del marcador del template."""
    run = paragraph.runs[0]
    font = run.font
    return (
        paragraph.style.name,
        paragraph.alignment,
        paragraph.paragraph_format.space_before,
        paragraph.paragraph_format.space_after,
        paragraph.paragraph_format.line_spacing,
        paragraph.paragraph_format.line_spacing_rule,
        run.style.name,
        font.name,
        font.size,
        font.bold,
        font.italic,
        font.color.type,
        font.color.rgb,
        font.underline,
        font.all_caps,
        font.small_caps,
        font.strike,
        font.double_strike,
    )


class ApiEvidenceGeneratorTests(unittest.TestCase):
    def test_parse_acceptance_criteria_groups_given_when_then_content(self) -> None:
        source = """CA01: Visualización inicial del módulo
Dado que el asesor ingresa al módulo de solicitudes,
Cuando se cargue la pantalla,
Entonces debe visualizar el título “Solicitudes” y las pestañas “Pre-solicitudes” y “Solicitudes”.

CA02: Estado vacío (empty state) Pre-solicitudes
Dado que el asesor no tiene pre-solicitudes disponibles,
Cuando ingrese a la pestaña Pre-solicitudes,
Entonces debe mostrarse un estado vacío indicando que no existen pre-solicitudes para gestionar."""

        cases = parse_acceptance_criteria(source)

        self.assertEqual(2, len(cases))
        self.assertEqual("CASO-001", cases[0].xray_test)
        self.assertEqual("CA01", cases[0].acceptance_criterion)
        self.assertEqual("Visualización inicial del módulo", cases[0].title)
        self.assertEqual(
            "Dado que el asesor ingresa al módulo de solicitudes,\nCuando se cargue la pantalla,\n"
            "Entonces debe visualizar el título “Solicitudes” y las pestañas “Pre-solicitudes” y “Solicitudes”.",
            cases[0].content,
        )
        self.assertEqual("CA02", cases[1].acceptance_criterion)
        self.assertEqual("Estado vacío (empty state) Pre-solicitudes", cases[1].title)
        self.assertIn("Dado que el asesor no tiene pre-solicitudes disponibles", cases[1].content)

    def test_parse_test_titles_uses_acceptance_criteria_as_cases(self) -> None:
        source = """CA01: Primer criterio
Dado algo
Cuando ocurre algo
Entonces resultado

CA02: Segundo criterio
Dado otra condición
Y otra condición
Cuando ocurre algo
Entonces resultado

CA03: Tercer criterio
Dado algo más
Cuando ocurre
Pero existe otra condición
Entonces resultado final"""

        cases = parse_test_titles(source)

        self.assertEqual(3, len(cases))
        self.assertEqual(["CA01", "CA02", "CA03"], [case.acceptance_criterion for case in cases])
        self.assertEqual(["CASO-001", "CASO-002", "CASO-003"], [case.xray_test for case in cases])
        self.assertIn("Y otra condición", cases[1].content)
        self.assertIn("Pero existe otra condición", cases[2].content)

    def test_parse_acceptance_criteria_supports_ca_variants(self) -> None:
        cases = parse_test_titles("CA1: Caso uno\n\nCA 02: Caso dos\n\nca03: Caso tres")

        self.assertEqual(3, len(cases))
        self.assertEqual(["CA1", "CA02", "CA03"], [case.acceptance_criterion for case in cases])

    def test_parse_test_titles_accepts_plain_jira_title_list(self) -> None:
        cases = parse_test_titles("Validar ingreso\nValidar cliente elegible\n")

        self.assertEqual(["CASO-001", "CASO-002"], [case.xray_test for case in cases])
        self.assertEqual(["Validar ingreso", "Validar cliente elegible"], [case.title for case in cases])

    def test_parse_xray_tests_preserves_case_title_and_order(self) -> None:
        source = """Xray Test
TAP-79281

T-ADN1-2694 Error por salesRepReference no resuelto

Xray Test
TAP-79280

T-ADN1-2694 Respuesta separada por secciones

Xray Test
TAP-79279

T-ADN1-2694 Consulta pre-solicitudes PRE_APPLICATION

Xray Test
TAP-79278

T-ADN1-2694 Consulta solicitudes APPLICATION"""

        cases = parse_xray_tests(source)

        self.assertEqual(["TAP-79281", "TAP-79280", "TAP-79279", "TAP-79278"], [case.xray_test for case in cases])
        self.assertEqual(
            "T-ADN1-2694 Error por salesRepReference no resuelto",
            cases[0].title,
        )
        self.assertEqual("POST", cases[0].method)

    def test_evidence_case_defaults_to_post(self) -> None:
        case = EvidenceCase(
            "TAP-79281",
            "Caso simplificado",
            endpoint="/legacy-endpoint",
            status_response="200",
            request='{"legacy": true}',
        )

        self.assertEqual("POST", case.method)

    def test_parse_xray_test_executions_format(self) -> None:
        source = """1
TAP-79276
T-ADN1-2694 Consulta bandeja de solicitudes del AdN
Cucumber
0
TO DO

2
TAP-79277
T-ADN1-2694 Cálculo automático del periodo
Cucumber
0
TO DO"""

        cases = parse_xray_tests(source)

        self.assertEqual(["TAP-79276", "TAP-79277"], [case.xray_test for case in cases])
        self.assertEqual(
            [
                "T-ADN1-2694 Consulta bandeja de solicitudes del AdN",
                "T-ADN1-2694 Cálculo automático del periodo",
            ],
            [case.title for case in cases],
        )

    def test_parse_xray_test_executions_keeps_six_cases_in_order(self) -> None:
        source = """1
TAP-79276
T-ADN1-2694 Consulta bandeja de solicitudes del AdN
Cucumber
0
TO DO

2
TAP-79277
T-ADN1-2694 Cálculo automático del periodo
Cucumber
0
TO DO

3
TAP-79278
T-ADN1-2694 Consulta solicitudes APPLICATION
Cucumber
0
TO DO

4
TAP-79279
T-ADN1-2694 Consulta pre-solicitudes PRE_APPLICATION
Cucumber
0
TO DO

5
TAP-79280
T-ADN1-2694 Respuesta separada por secciones
Cucumber
0
TO DO

6
TAP-79281
T-ADN1-2694 Error por salesRepReference no resuelto
Cucumber
0
TO DO"""

        cases = parse_xray_tests(source)

        self.assertEqual(
            ["TAP-79276", "TAP-79277", "TAP-79278", "TAP-79279", "TAP-79280", "TAP-79281"],
            [case.xray_test for case in cases],
        )
        self.assertEqual("T-ADN1-2694 Error por salesRepReference no resuelto", cases[-1].title)

    def test_parse_xray_tests_supports_tabs_and_ignores_metadata(self) -> None:
        source = (
            "1\tTAP-80001\tT-ABC-1234 Primer caso\tCucumber\t0\tTO DO\n"
            "2\tTAP-80002\tT-XYZ-999 Segundo caso\tManual\t0\tPASS\n"
            "3\tTAP-80003\tT-ABC-1235 Tercer caso\tCucumber\t0\tBLOCKED"
        )

        cases = parse_xray_tests(source)

        self.assertEqual(["TAP-80001", "TAP-80002", "TAP-80003"], [case.xray_test for case in cases])
        self.assertEqual(
            ["T-ABC-1234 Primer caso", "T-XYZ-999 Segundo caso", "T-ABC-1235 Tercer caso"],
            [case.title for case in cases],
        )
        metadata = {"Cucumber", "0", "TO DO", "PASS", "FAIL", "BLOCKED"}
        self.assertTrue(all(case.title not in metadata for case in cases))

    def test_parse_real_jira_markdown_table(self) -> None:
        source = """| 1  | [TAP-79276](https://mibancoinnova.atlassian.net/browse/TAP-79276) | T-ADN1-2694 Consulta bandeja de solicitudes del AdN | Cucumber | | 0 | **TO DO** |
| :- | :---------------------------------------------------------------- | :-------------------------------------------------- | :------- | :- | - | :-------- |
|    | | 2 | [TAP-79277](https://mibancoinnova.atlassian.net/browse/TAP-79277) | T-ADN1-2694 Cálculo automático del periodo | Cucumber | | 0 | **TO DO** |
|    | | 3 | [TAP-79278](https://mibancoinnova.atlassian.net/browse/TAP-79278) | T-ADN1-2694 Consulta solicitudes APPLICATION | Cucumber | | 0 | **TO DO** |
|    | | 4 | [TAP-79279](https://mibancoinnova.atlassian.net/browse/TAP-79279) | T-ADN1-2694 Consulta pre-solicitudes PRE_APPLICATION | Cucumber | | 0 | **TO DO** |
|    | | 5 | [TAP-79280](https://mibancoinnova.atlassian.net/browse/TAP-79280) | T-ADN1-2694 Respuesta separada por secciones | Cucumber | | 0 | **TO DO** |
|    | | 6 | [TAP-79281](https://mibancoinnova.atlassian.net/browse/TAP-79281) | T-ADN1-2694 Error por salesRepReference no resuelto | Cucumber | | 0 | **TO DO** |"""

        cases = parse_xray_tests(source)

        self.assertEqual(
            ["TAP-79276", "TAP-79277", "TAP-79278", "TAP-79279", "TAP-79280", "TAP-79281"],
            [case.xray_test for case in cases],
        )
        self.assertEqual(
            [
                "T-ADN1-2694 Consulta bandeja de solicitudes del AdN",
                "T-ADN1-2694 Cálculo automático del periodo",
                "T-ADN1-2694 Consulta solicitudes APPLICATION",
                "T-ADN1-2694 Consulta pre-solicitudes PRE_APPLICATION",
                "T-ADN1-2694 Respuesta separada por secciones",
                "T-ADN1-2694 Error por salesRepReference no resuelto",
            ],
            [case.title for case in cases],
        )
        forbidden = ("Cucumber", "0", "TO DO", "**", "https://")
        self.assertTrue(all(not any(value in case.title for value in forbidden) for case in cases))

    def test_parse_markdown_handles_empty_columns_and_tap_duplicates(self) -> None:
        source = """| | | 6 | [TAP-79281](https://example.test/TAP-79281) | T-CUSTOMER-123 Consulta cliente | Cucumber | | 0 | **TO DO** |
| | 7 | [TAP-79281](https://example.test/TAP-79281) | T-CUSTOMER-999 Debe ignorarse por duplicado | Cucumber | | 0 | **PASS** |"""

        cases = parse_xray_tests(source)

        self.assertEqual(1, len(cases))
        self.assertEqual("TAP-79281", cases[0].xray_test)
        self.assertEqual("T-CUSTOMER-123 Consulta cliente", cases[0].title)

    def test_generates_one_populated_block_per_xray_test(self) -> None:
        template = resource_path("templates/Evidencias_APIs_template.docx")
        original_hash = sha256(template.read_bytes()).hexdigest()
        cases = [
            EvidenceCase("TAP-79281", "T-ADN1-2694 Error por salesRepReference no resuelto", "POST", "/sales", "400", "Bad Request", '{"id": 1}', '{"error": "unresolved"}', "Se valida el error."),
            EvidenceCase("TAP-79280", "T-ADN1-2694 Respuesta separada por secciones", "GET", "/sections", "200", "OK", "", '{"items": []}', "Se valida la respuesta."),
            EvidenceCase("TAP-79279", "T-ADN1-2694 Consulta pre-solicitudes PRE_APPLICATION"),
            EvidenceCase("TAP-79278", "T-ADN1-2694 Consulta solicitudes APPLICATION", "DELETE", "/requests/1"),
        ]
        data = ApiEvidenceData("ADN1-2694", "Diego Molina", "21/09/2026", "Customer Request", "TP-123", cases)

        with TemporaryDirectory() as temporary_directory:
            output = Path(temporary_directory) / evidence_filename(data.jira)
            generate_api_evidence(data, template, output)
            document = Document(output)

            general, *evidence_tables = document.tables
            self.assertEqual(5, len(document.tables))
            self.assertEqual(data.qa_agile, general.cell(0, 1).text)
            self.assertEqual("", general.cell(2, 1).text)
            self.assertEqual(DEFAULT_DOCUMENT_STATUS, general.cell(3, 1).text)
            self.assertEqual("", general.cell(3, 3).text)
            self.assertEqual(data.test_plan, general.cell(4, 1).text)
            self.assertEqual([case.xray_test for case in cases], [table.cell(0, 1).text for table in evidence_tables])
            self.assertEqual(cases[0].method, evidence_tables[0].cell(0, 3).text)
            self.assertEqual(cases[1].response, evidence_tables[1].cell(4, 1).text)
            self.assertEqual("POST", evidence_tables[2].cell(0, 3).text)
            self.assertEqual([f"Escenario de prueba {number}: {case.title}" for number, case in enumerate(cases, start=1)], [
                paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip().startswith("Escenario de prueba ")
            ])
            self.assertEqual([(6, 4)] * 4, [(len(table.rows), len(table.columns)) for table in evidence_tables])

            with ZipFile(template) as source, ZipFile(output) as generated:
                self.assertEqual(source.read("word/media/image1.png"), generated.read("word/media/image1.png"))

        self.assertEqual(original_hash, sha256(template.read_bytes()).hexdigest())

    def test_writes_selected_test_environment(self) -> None:
        template = resource_path("templates/Evidencias_APIs_template.docx")
        cases = [EvidenceCase("TAP-1", "Caso de entorno")]

        for environment in ("QA", "STG"):
            with self.subTest(environment=environment), TemporaryDirectory() as temporary_directory:
                data = ApiEvidenceData("ADN-1", "Diego", "23/09/2026", "/service", "TP-1", cases, environment)
                output = Path(temporary_directory) / evidence_filename(data.jira)
                generate_api_evidence(data, template, output)
                self.assertEqual(environment, Document(output).tables[0].cell(3, 3).text)

    def test_separates_api_domain_from_evidence_titles_using_last_slash(self) -> None:
        template = resource_path("templates/Evidencias_APIs_template.docx")
        domain = "Urpipro/Loan/Derivacion App MiBanco"
        cases = [
            EvidenceCase("TAP-78744", f"{domain}/Visualización correcta del modal de confirmación"),
            EvidenceCase("TAP-78745", f"{domain}/Confirmación y procesamiento de la solicitud"),
            EvidenceCase("TAP-78746", f"{domain}/Cancelación de la operación"),
        ]
        data = ApiEvidenceData("ADN-234", "Diego", "23/09/2026", "/service", "TP-1", cases, "QA")

        with TemporaryDirectory() as temporary_directory:
            output = Path(temporary_directory) / evidence_filename(data.jira)
            generate_api_evidence(data, template, output)
            document = Document(output)

        self.assertEqual(domain, document.tables[0].cell(2, 1).text)
        self.assertEqual(
            [
                "Escenario de prueba 1: Visualización correcta del modal de confirmación",
                "Escenario de prueba 2: Confirmación y procesamiento de la solicitud",
                "Escenario de prueba 3: Cancelación de la operación",
            ],
            [paragraph.text for paragraph in document.paragraphs if paragraph.text.startswith("Escenario de prueba ")],
        )

    def test_uses_complete_title_as_evidence_when_api_title_has_no_slash(self) -> None:
        template = resource_path("templates/Evidencias_APIs_template.docx")
        data = ApiEvidenceData("ADN-234", "Diego", "23/09/2026", "/service", "TP-1", [EvidenceCase("TAP-1", "Sin ruta")])

        with TemporaryDirectory() as temporary_directory:
            output = Path(temporary_directory) / evidence_filename(data.jira)
            generate_api_evidence(data, template, output)
            document = Document(output)

        self.assertEqual("", document.tables[0].cell(2, 1).text)
        self.assertIn("Escenario de prueba 1: Sin ruta", [paragraph.text for paragraph in document.paragraphs])

    def test_api_scenario_title_preserves_the_mobile_template_format(self) -> None:
        api_template = resource_path("templates/Evidencias_APIs_template.docx")
        mobile_template = resource_path("templates/Evidencias_Mobile_template.docx")
        mobile_heading = next(
            paragraph
            for paragraph in Document(mobile_template).paragraphs
            if paragraph.text.strip().startswith("Escenario de prueba:")
        )
        data = ApiEvidenceData(
            "ADN-234", "Diego", "23/09/2026", "/service", "TP-1", [EvidenceCase("TAP-1", "Un escenario")]
        )

        with TemporaryDirectory() as temporary_directory:
            output = Path(temporary_directory) / evidence_filename(data.jira)
            generate_api_evidence(data, api_template, output)
            api_heading = next(
                paragraph
                for paragraph in Document(output).paragraphs
                if paragraph.text.startswith("Escenario de prueba 1:")
            )

        self.assertEqual(_scenario_title_format(mobile_heading), _scenario_title_format(api_heading))


if __name__ == "__main__":
    unittest.main()
