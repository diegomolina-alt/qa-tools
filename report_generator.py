"""Generación del informe DOCX de pruebas."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.shared import Pt

from qa_engine import TestCase


def generate_report(
    destination: Path,
    jira: str,
    analyst: str,
    criteria_text: str,
    test_cases: list[TestCase],
) -> None:
    """Crea un informe limpio con la información y los casos generados."""
    document = Document()
    normal_style = document.styles["Normal"]
    normal_style.font.name = "Arial"
    normal_style.font.size = Pt(10)

    document.add_heading("INFORME DE PRUEBAS", level=0)
    document.add_heading("Información general", level=1)
    document.add_paragraph(f"Número Jira: {jira}")
    document.add_paragraph(f"Analista QA: {analyst}")

    document.add_heading("CRITERIOS DE ACEPTACIÓN", level=1)
    document.add_paragraph(criteria_text)

    document.add_heading("CASOS DE PRUEBA", level=1)
    for case in test_cases:
        document.add_heading(case.id, level=2)
        _add_field(document, "Criterio", case.criterio)
        _add_field(document, "Título", case.titulo)
        _add_field(document, "Precondiciones", case.precondiciones)
        _add_field(document, "Pasos", case.pasos)
        _add_field(document, "Resultado esperado", case.resultado_esperado)
        _add_field(document, "Tipo", case.tipo)

    document.save(destination)


def _add_field(document: Document, label: str, value: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.add_run(f"{label}: ").bold = True
    paragraph.add_run(value)
