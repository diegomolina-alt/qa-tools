"""Parsing y generación del encabezado de evidencias Mobile desde su plantilla."""

from __future__ import annotations

from copy import deepcopy
from dataclasses import dataclass, replace
from datetime import datetime
from pathlib import Path
import re
from typing import Mapping

from docx import Document
from docx.oxml import OxmlElement
from docx.table import _Cell
from docx.text.paragraph import Paragraph

from docx_helpers import set_scenario_title_text
from filename_helpers import sanitize_filename_component


REQUIRED_FIELDS = {
    "QA_AGILE": "QA Agile",
    "FECHA_EJECUCION": "Fecha de ejecución",
    "CASO_PRUEBA": "Caso de Prueba",
    "ESTADO": "Estado",
    "ENTORNO_PRUEBA": "Entorno de prueba",
}
ALLOWED_STATUSES = ("PASSED", "FAILED", "BLOCKED", "NOT_EXECUTED")
MOBILE_EXPECTED_RESULT = (
    "Que se cumplan todos los escenarios planificados. Todos los escenarios deben finalizar satisfactoriamente y cumplir con lo esperado."
)
TEMPLATE_LABELS = {
    "QA Agile:": "QA_AGILE",
    "Fecha Ejecución:": "FECHA_EJECUCION",
    "Caso de Prueba:": "CASO_PRUEBA",
    # La versión vigente de la plantilla usa Dominio para este mismo valor.
    "Dominio:": "CASO_PRUEBA",
    "Estado:": "ESTADO",
    "Precondiciones:": "PRECONDICIONES",
    "Entorno de prueba:": "ENTORNO_PRUEBA",
    "Resultado Esperado:": "RESULTADO_ESPERADO",
    "Test Plan:": "TEST_PLAN",
}


@dataclass(frozen=True)
class MobileEvidenceData:
    qa_agile: str = ""
    execution_date: str = ""
    test_case: str = ""
    status: str = ""
    preconditions: str = ""
    test_environment: str = ""
    expected_result: str = ""
    test_plan: str = ""

    @classmethod
    def from_mapping(cls, values: Mapping[str, str]) -> "MobileEvidenceData":
        return cls(
            qa_agile=values.get("QA_AGILE", "").strip(),
            execution_date=values.get("FECHA_EJECUCION", "").strip(),
            test_case=values.get("CASO_PRUEBA", "").strip(),
            status=values.get("ESTADO", "").strip(),
            preconditions=values.get("PRECONDICIONES", "").strip(),
            test_environment=values.get("ENTORNO_PRUEBA", "").strip(),
            expected_result=values.get("RESULTADO_ESPERADO", "").strip(),
            test_plan=values.get("TEST_PLAN", "").strip(),
        )

    def as_template_values(self) -> dict[str, str]:
        return {
            "QA_AGILE": self.qa_agile,
            "FECHA_EJECUCION": self.execution_date,
            "CASO_PRUEBA": self.test_case,
            "ESTADO": self.status,
            "PRECONDICIONES": self.preconditions,
            "ENTORNO_PRUEBA": self.test_environment,
            "RESULTADO_ESPERADO": self.expected_result,
            "TEST_PLAN": self.test_plan,
        }


def parse_txt(text: str) -> dict[str, str]:
    """Lee pares CLAVE=valor, conservando UTF-8 y separando solo el primer '='."""
    values: dict[str, str] = {}
    for line_number, raw_line in enumerate(text.splitlines(), start=1):
        line = raw_line.strip()
        if not line:
            continue
        if "=" not in line:
            raise ValueError(f"Línea {line_number} inválida. Utilice el formato CLAVE=valor.")
        key, value = line.split("=", 1)
        key = key.strip()
        if not key:
            raise ValueError(f"Línea {line_number} inválida. La clave no puede estar vacía.")
        values[key] = value.strip()
    return values


def validate_report_data(data: MobileEvidenceData) -> list[str]:
    """Devuelve todos los errores de la cabecera para presentarlos en la interfaz."""
    values = data.as_template_values()
    errors = [f"El campo {label} es obligatorio" for key, label in REQUIRED_FIELDS.items() if not values[key]]

    if data.execution_date:
        try:
            datetime.strptime(data.execution_date, "%d/%m/%Y")
        except ValueError:
            errors.append("Fecha de ejecución inválida. Utilice DD/MM/YYYY")

    if data.status and data.status not in ALLOWED_STATUSES:
        errors.append("Estado inválido. Estados permitidos: PASSED, FAILED, BLOCKED, NOT_EXECUTED")
    return errors


def load_mobile_template(template_path: Path) -> Document:
    """Abre la plantilla aprobada sin modificarla."""
    document = Document(template_path)
    if len(document.tables) < 2:
        raise ValueError("La plantilla Mobile no contiene la cabecera y el bloque de escenario requeridos.")

    found_keys = {
        key
        for table in document.tables
        for row in table.rows
        for cell in row.cells
        if (key := _template_key_for_cell(cell.text)) is not None
    }
    if set(TEMPLATE_LABELS.values()) - found_keys:
        raise ValueError("La plantilla Mobile no contiene las etiquetas de cabecera requeridas.")

    scenario_heading = document.tables[1]._tbl.getprevious()
    if scenario_heading is None or "Escenario de prueba:" not in Paragraph(scenario_heading, document._body).text:
        raise ValueError("La plantilla Mobile no contiene el bloque 'Escenario de prueba:' requerido.")
    return document


def fill_mobile_header(document: Document, data: MobileEvidenceData, *, preserve_empty_values: bool = False) -> None:
    """Completa las celdas ubicadas por etiqueta, preservando el formato de la plantilla."""
    values = data.as_template_values()
    found_keys: set[str] = set()
    for table in document.tables:
        for row_index, row in enumerate(table.rows):
            for column_index, cell in enumerate(row.cells):
                key = _template_key_for_cell(cell.text)
                if key is None:
                    continue
                # "Estado:" también aparece en la sección de escenarios, que aún
                # no forma parte de esta iteración. Solo se completa la primera
                # coincidencia de cada etiqueta de la cabecera.
                if key in found_keys:
                    continue
                if column_index + 1 >= len(row.cells):
                    raise ValueError(f"La etiqueta {_normalized_label(cell.text)} no tiene una celda de valor.")
                if values[key] or not preserve_empty_values:
                    _write_cell_text(table.cell(row_index, column_index + 1), values[key])
                found_keys.add(key)

    missing = set(TEMPLATE_LABELS.values()) - found_keys
    if missing:
        raise ValueError("La plantilla Mobile no contiene todas las etiquetas requeridas.")


def generate_mobile_report(data: MobileEvidenceData, template_path: Path, output_path: Path) -> None:
    """Genera un nuevo DOCX a partir de la plantilla, dejando el original intacto."""
    errors = validate_report_data(data)
    if errors:
        raise ValueError("\n".join(errors))
    document = load_mobile_template(template_path)
    fill_mobile_header(document, data)
    document.save(output_path)


def generate_mobile_evidence(
    data: MobileEvidenceData, cases: list[object], template_path: Path, output_path: Path
) -> None:
    """Genera una cabecera parcial y un bloque de escenario clonado por cada caso detectado."""
    if not cases:
        raise ValueError("Debe existir al menos un caso de evidencia.")

    case_path = ""
    scenario_names: list[str] = []
    for case in cases:
        detected_case_path, scenario_name = _split_mobile_title(str(getattr(case, "title", "")))
        if detected_case_path and not case_path:
            case_path = detected_case_path
        scenario_names.append(scenario_name)

    document = load_mobile_template(template_path)
    # Los campos sin fuente fiable conservan el valor y estilo aprobados de la plantilla.
    fill_mobile_header(
        document,
        replace(data, test_case=case_path, expected_result=MOBILE_EXPECTED_RESULT) if case_path else replace(data, expected_result=MOBILE_EXPECTED_RESULT),
        preserve_empty_values=True,
    )
    scenario_table = document.tables[1]
    scenario_heading = scenario_table._tbl.getprevious()
    blocks: list[tuple[object, _Cell]] = [(scenario_heading, scenario_table)]
    for _ in cases[1:]:
        blocks.append(_duplicate_scenario_block(blocks[-1][1], document, scenario_heading))

    for number, (case, scenario_name, (heading, table)) in enumerate(zip(cases, scenario_names, blocks), start=1):
        set_scenario_title_text(
            Paragraph(heading, document._body), f"Escenario de prueba {number}: {scenario_name}"
        )
        _write_case_identifier(table.cell(0, 0), str(getattr(case, "xray_test", "")))
        _write_cell_text(table.cell(0, 1), "Estado: Validado")

    document.save(output_path)


def mobile_evidence_filename(jira: str) -> str:
    """Crea Evidencia_Mobile_<JIRA>.docx sin alterar guiones válidos del Jira."""
    safe_jira = sanitize_filename_component(jira, fallback="SIN_JIRA")
    return f"Evidencia_Mobile_{safe_jira}.docx"


def _normalized_label(value: str) -> str:
    return " ".join(value.split())


def _template_key_for_cell(value: str) -> str | None:
    """Reconoce una etiqueta aunque la plantilla añada una aclaración en otra línea."""
    normalized = _normalized_label(value)
    for label, key in TEMPLATE_LABELS.items():
        if normalized.startswith(label):
            return key
    return None


def _write_cell_text(cell: _Cell, value: str) -> None:
    """Reutiliza el primer run y párrafo de la celda para conservar su estilo visual."""
    paragraph: Paragraph = cell.paragraphs[0]
    if paragraph.runs:
        paragraph.runs[0].text = value
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(value)
    for extra_paragraph in cell.paragraphs[1:]:
        for run in extra_paragraph.runs:
            run.text = ""


def _write_case_identifier(cell: _Cell, case_id: str) -> None:
    """Conserva el formato del rótulo Caso: y añade únicamente el identificador real."""
    paragraph: Paragraph = cell.paragraphs[0]
    if paragraph.runs:
        paragraph.runs[0].text = "Caso: "
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run("Caso: ")
    paragraph.add_run(case_id)


def _duplicate_scenario_block(table: object, document: Document, template_heading: object) -> tuple[object, object]:
    """Clona el párrafo y tabla del escenario, preservando su XML de formato."""
    separator = OxmlElement("w:p")
    copied_heading = deepcopy(template_heading)
    copied_table = deepcopy(table._tbl)
    table._tbl.addnext(separator)
    separator.addnext(copied_heading)
    copied_heading.addnext(copied_table)
    from docx.table import Table

    return copied_heading, Table(copied_table, document._body)


def _split_mobile_title(title: str) -> tuple[str, str]:
    """Separa la ruta de caso del escenario por la última barra disponible."""
    if "/" not in title:
        return "", title.strip()
    case_path, scenario_name = title.rsplit("/", 1)
    return case_path.strip(), scenario_name.strip()
